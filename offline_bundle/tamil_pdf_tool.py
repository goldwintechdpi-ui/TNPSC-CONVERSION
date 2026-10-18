"""
tamil_pdf_tool.py - ONE FILE: batch-convert Tamil/English PDFs (or images) with ocr_tamil and
extract the numbered questions (default 1 to 100) into Excel / text / Word.

  * Select PDF files or a whole folder in the window; results are saved in the same folder,
    with the same file name.
  * "Extract questions" mode: Question column = only the question text, Anwers column = only the
    options (A)-(E). Junk characters are removed. Numbers outside the chosen range are dropped.
  * Built on the library's own steps (CRAFT word detection + PARSeq reading), reading the word crops
    of several pages in one pooled batch, with column-aware reading order.

Install:
    pip install ocr_tamil pymupdf openpyxl python-docx
    (optional, better junk detection: pip install pyspellchecker)

Run:
    python tamil_pdf_tool.py
"""
import difflib
import inspect
import queue
import re
import sys
import threading
import time
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np


# ============================================================================
# OCR + PAGES (batch pipeline built on the ocr_tamil library)
# ============================================================================
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def parse_pages(spec, total):
    """'1-3,5' -> [0,1,2,4] (zero-based). Empty -> all pages."""
    if not spec:
        return list(range(total))
    pages = set()
    for part in str(spec).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-")
            pages.update(range(int(a) - 1, min(int(b), total)))
        else:
            pages.add(int(part) - 1)
    return sorted(p for p in pages if 0 <= p < total)


def count_pages(path, pages_spec=None):
    """Number of pages that will be read (1 for images, 1 if a PDF cannot be opened)."""
    if str(path).lower().endswith(".pdf"):
        try:
            import fitz

            with fitz.open(str(path)) as d:
                return max(1, len(parse_pages(pages_spec, len(d))))
        except Exception:
            return 1
    return 1


def read_image_file(path):
    """Read an image as a BGR array. Works with Tamil / non-ASCII file names on Windows."""
    import cv2

    data = np.fromfile(str(path), dtype=np.uint8)
    img = cv2.imdecode(data, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError(f"Cannot read image: {path}")
    return img


def iter_pdf_pages(pdf_path, dpi=150, pages_spec=None):
    """Yield (page_number, BGR image array), rendering one page at a time in memory."""
    import fitz  # PyMuPDF

    zoom = dpi / 72.0
    with fitz.open(str(pdf_path)) as doc:
        for i in parse_pages(pages_spec, len(doc)):
            pix = doc[i].get_pixmap(matrix=fitz.Matrix(zoom, zoom), colorspace=fitz.csRGB, alpha=False)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            yield i + 1, np.ascontiguousarray(arr[:, :, ::-1])


def load_engine(langs=("tamil", "english"), batch_size=32, min_conf=0.85, use_gpu=True, straight_page=True):
    from ocr_tamil.ocr import OCR

    wanted = dict(
        detect=True, details=2, lang=list(langs), batch_size=int(batch_size),
        recognize_thres=float(min_conf), enable_cuda=bool(use_gpu), assume_straight_page=bool(straight_page),
    )
    supported = set(inspect.signature(OCR.__init__).parameters)
    return OCR(**{k: v for k, v in wanted.items() if k in supported})


def engine_device(ocr):
    return str(getattr(ocr, "device", "cpu"))


def box_xyxy(box):
    arr = np.asarray(box, dtype=float).reshape(-1, 2)
    return (float(arr[:, 0].min()), float(arr[:, 1].min()), float(arr[:, 0].max()), float(arr[:, 1].max()))


def _read_crops(ocr, crops):
    texts = [""] * len(crops)
    confs = [0.0] * len(crops)
    good = [i for i, c in enumerate(crops)
            if c is not None and getattr(c, "ndim", 0) >= 2 and c.shape[0] >= 2 and c.shape[1] >= 2]
    if not good:
        return texts, confs
    t, c = ocr.text_recognize_batch([crops[i] for i in good])
    for i, tt, cc in zip(good, t, c):
        texts[i], confs[i] = tt, float(cc)
    return texts, confs


def find_gutters(boxes, page_w, min_gap=0.025):
    n = len(boxes)
    if n < 12 or page_w <= 0:
        return []
    bins = 400
    cover = np.zeros(bins + 1, dtype=int)
    lo = np.clip((boxes[:, 0] / page_w * bins).astype(int), 0, bins)
    hi = np.clip((boxes[:, 2] / page_w * bins).astype(int), 0, bins)
    for a, b in zip(lo, hi):
        cover[a:b + 1] += 1
    tol = max(1, int(0.01 * n))
    min_w = max(3, int(min_gap * bins))
    centers = (boxes[:, 0] + boxes[:, 2]) / 2
    gutters, i, end = [], int(0.15 * bins), int(0.85 * bins)
    while i <= end:
        if cover[i] <= tol:
            j = i
            while j <= end and cover[j] <= tol:
                j += 1
            if j - i >= min_w:
                x = ((i + j - 1) / 2) / bins * page_w
                left = int((centers < x).sum())
                if left >= 0.2 * n and (n - left) >= 0.2 * n:
                    gutters.append(x)
            i = j
        else:
            i += 1
    return gutters


def arrange_words(words, page_w, columns=True):
    if not words:
        return []
    boxes = np.array([w["xyxy"] for w in words], dtype=float)
    splits = find_gutters(boxes, page_w) if columns else []
    cx = (boxes[:, 0] + boxes[:, 2]) / 2
    col_of = np.searchsorted(np.array(splits), cx) if splits else np.zeros(len(words), dtype=int)

    ordered, line_no = [], 0
    for c in range(len(splits) + 1):
        idx = [i for i in range(len(words)) if col_of[i] == c]
        if not idx:
            continue
        mh = float(np.median(boxes[idx, 3] - boxes[idx, 1])) or 1.0
        idx.sort(key=lambda i: (boxes[i, 1] + boxes[i, 3]) / 2)
        lines, cur, cur_y = [], [], 0.0
        for i in idx:
            yc = (boxes[i, 1] + boxes[i, 3]) / 2
            if cur and abs(yc - cur_y) > 0.6 * mh:
                lines.append(cur)
                cur = []
            cur.append(i)
            cur_y = float(np.mean([(boxes[k, 1] + boxes[k, 3]) / 2 for k in cur]))
        if cur:
            lines.append(cur)
        for ln in lines:
            line_no += 1
            for i in sorted(ln, key=lambda k: boxes[k, 0]):
                w = dict(words[i])
                w["line"], w["col"] = line_no, c + 1
                w["bbox"] = tuple(round(v, 1) for v in boxes[i])
                ordered.append(w)
    return ordered


def recognize_pages(ocr, chunk, columns=True):
    if not (hasattr(ocr, "craft_detect") and hasattr(ocr, "text_recognize_batch")):
        return _recognize_with_predict(ocr, chunk, columns)

    t0 = time.time()
    per_page, pool = [], []
    for page_no, img in chunk:
        crops, preds = ocr.craft_detect(img)
        per_page.append((page_no, img.shape[1], preds, len(pool), len(crops)))
        pool.extend(crops)
    t_det = time.time() - t0

    t1 = time.time()
    texts, confs = _read_crops(ocr, pool)
    t_read = time.time() - t1

    out = {}
    for page_no, width, preds, start, n in per_page:
        words = []
        for k in range(n):
            text = (texts[start + k] or "").strip()
            if not text:
                continue
            words.append({"text": text, "conf": float(confs[start + k]), "xyxy": box_xyxy(preds[k][0])})
        out[page_no] = arrange_words(words, width, columns)
    return out, t_det, t_read


def _recognize_with_predict(ocr, chunk, columns):
    t0 = time.time()
    out = {}
    for page_no, img in chunk:
        words = []
        for item in (ocr.predict([img]) or [[]])[0]:
            text = str(item[0]).strip()
            if text:
                words.append({"text": text, "conf": float(item[1]), "xyxy": box_xyxy(item[2][0])})
        out[page_no] = arrange_words(words, img.shape[1], columns)
    return out, 0.0, time.time() - t0


def ocr_document(ocr, src, dpi=150, pages_spec=None, page_batch=4, columns=True,
                 should_stop=None, on_progress=None):
    src = Path(src)
    if src.suffix.lower() == ".pdf":
        pages = iter_pdf_pages(src, dpi, pages_spec)
    else:
        pages = iter([(1, read_image_file(src))])

    results, chunk = {}, []

    def flush():
        t = time.time()
        out, t_det, t_read = recognize_pages(ocr, chunk, columns)
        results.update(out)
        if on_progress:
            on_progress({"pages": [p for p, _ in chunk], "seconds": time.time() - t,
                         "detect": t_det, "read": t_read,
                         "words": sum(len(v) for v in out.values())})
        chunk.clear()

    for item in pages:
        if should_stop and should_stop():
            return results, True
        chunk.append(item)
        if len(chunk) >= max(1, int(page_batch)):
            flush()
    if chunk:
        if should_stop and should_stop():
            return results, True
        flush()
    return results, False


def group_lines(words):
    lines = {}
    for w in words:
        lines.setdefault(w["line"], []).append(w)
    out = []
    for line_no in sorted(lines):
        ws = lines[line_no]
        out.append((line_no, ws[0].get("col", 1), " ".join(w["text"] for w in ws),
                    sum(w["conf"] for w in ws) / len(ws)))
    return out


def flat_bbox(bbox):
    try:
        if bbox is None:
            return ("", "", "", "")
        if hasattr(bbox, "tolist"):
            bbox = bbox.tolist()
        if isinstance(bbox[0], (list, tuple)):
            xs = [p[0] for p in bbox]
            ys = [p[1] for p in bbox]
            return (min(xs), min(ys), max(xs), max(ys))
        return tuple(bbox[:4])
    except Exception:
        return ("", "", "", "")


def save_txt(results, path):
    with open(path, "w", encoding="utf-8") as f:
        for page_no, words in results.items():
            f.write(f"===== Page {page_no} =====\n")
            for _, _, text, _ in group_lines(words):
                f.write(text + "\n")
            f.write("\n")


def save_docx(results, path):
    from docx import Document

    doc = Document()
    first = True
    for page_no, words in results.items():
        if not first:
            doc.add_page_break()
        first = False
        doc.add_heading(f"Page {page_no}", level=2)
        for _, _, text, _ in group_lines(words):
            doc.add_paragraph(text)
    doc.save(path)


def save_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()

    ws = wb.active
    ws.title = "Lines"
    ws.append(["Page", "Column", "Line", "Text", "Avg confidence"])
    for page_no, words in results.items():
        for line_no, col, text, conf in group_lines(words):
            ws.append([page_no, col, line_no, text, round(conf, 3)])

    wd = wb.create_sheet("Words")
    wd.append(["Page", "Column", "Line", "Word", "Confidence", "x1", "y1", "x2", "y2"])
    for page_no, words in results.items():
        for w in words:
            wd.append([page_no, w.get("col", 1), w["line"], w["text"], round(w["conf"], 3),
                       *flat_bbox(w.get("bbox"))])

    for sheet in (ws, wd):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    ws.column_dimensions["D"].width = 80
    wd.column_dimensions["D"].width = 30
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


SAVERS = {"xlsx": save_xlsx, "docx": save_docx, "txt": save_txt}


def collect_inputs(target, recursive):
    target = Path(target)
    if target.is_file():
        return [target]
    if not target.is_dir():
        sys.exit(f"Not found: {target}")
    pattern = "**/*" if recursive else "*"
    return [
        p for p in sorted(target.glob(pattern))
        if p.is_file() and (p.suffix.lower() == ".pdf" or p.suffix.lower() in IMAGE_EXTS)
    ]


def output_stem(src, all_files):
    clash = [f for f in all_files if f.parent == src.parent and f.stem == src.stem and f != src]
    if clash:
        return src.parent / f"{src.stem}_{src.suffix.lower().lstrip('.')}"
    return src.parent / src.stem


def fmt_time(sec):
    sec = int(max(0, sec))
    h, r = divmod(sec, 3600)
    m, s = divmod(r, 60)
    return f"{h}h {m:02d}m" if h else (f"{m}m {s:02d}s" if m else f"{s}s")


# ============================================================================
# TEXT CLEANING
# ============================================================================
TAMIL = "஀-௿"
E_CANON = "விடை தெரியவில்லை"

try:
    from spellchecker import SpellChecker

    _SP = SpellChecker()

    def _is_word(w):
        return w.lower() in _SP
except Exception:
    _SP = None

    def _is_word(w):
        w = w.lower()
        if len(w) < 3 or not w.isalpha():
            return False
        vowels = sum(c in "aeiouy" for c in w)
        if not (0.28 <= vowels / len(w) <= 0.65):
            return False
        if re.search(r"[^aeiouy]{4,}", w):
            return False
        if re.search(r"(.)\1\1", w):
            return False
        return True


_INVISIBLE = re.compile("[​-‏‪-‮⁠﻿­]")
_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_BAD_CHARS = re.compile(rf"[^{TAMIL}A-Za-z0-9\s.,;:!?'\"‘’“”()\[\]\-–—/…=+]")
_TAMIL_RE = re.compile(f"[{TAMIL}]")
_LATIN_RE = re.compile("[A-Za-z]")
_CORE_RE = re.compile(rf"^[^{TAMIL}A-Za-z0-9]+|[^{TAMIL}A-Za-z0-9]+$")
_ELLIPSIS_RE = re.compile(r"^[.…]{3,}$")
_LABEL_TOKEN = re.compile(r"^[\(\[][A-Ea-eRr][\)\]][.,:;]?$")
_PUNCT_KEEP = {":", ";", "-", "–", "—", "=", "?", "'"}


def normalize(text):
    text = unicodedata.normalize("NFC", str(text))
    text = _INVISIBLE.sub("", text)
    text = _CONTROL.sub("", text)
    text = text.replace(" ", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = _BAD_CHARS.sub(" ", text)
    return text


def classify(tok, answer_mode=False):
    if _ELLIPSIS_RE.match(tok):
        return "G", tok
    if _LABEL_TOKEN.match(tok):
        return "G", tok
    core = _CORE_RE.sub("", tok)
    if not core:
        return ("P", tok) if tok in _PUNCT_KEEP else ("J", "")
    has_t = bool(_TAMIL_RE.search(core))
    has_l = bool(_LATIN_RE.search(core))
    if has_t:
        if has_l:
            tok = _LATIN_RE.sub("", tok)
            core = _CORE_RE.sub("", tok)
            if not core:
                return "J", ""
        digits_free = re.sub(r"[0-9]", "", core)
        if re.search(r"[0-9]", core) and len(digits_free) <= 2:
            return "J", ""
        tamil_len = len(_TAMIL_RE.findall(core))
        if tok == core and tamil_len <= 2:
            return "T", tok
        return "G", tok
    if re.fullmatch(r"[0-9]+", core):
        return "N", tok
    if has_l:
        if re.search(r"[0-9]", core):
            return "J", ""
        if len(core) == 1:
            return ("G", tok) if core == "X" else ("J", "")
        if answer_mode and _SP is None:
            return "J", ""
        return ("G", tok) if _is_word(core) else ("J", "")
    return "J", ""


def clean_line(line, answer_mode=False):
    toks = line.split()
    if not toks:
        return "", 0, 0
    info = [classify(t, answer_mode) for t in toks]
    kinds = [k for k, _ in info]

    if answer_mode and all(k in ("T", "P", "N") for k in kinds):
        keep = [t for (k, t) in info if k != "J"]
        return " ".join(keep), len(toks), len(toks) - len(keep)

    def neighbor_kind(i, step):
        j = i + step
        while 0 <= j < len(kinds):
            if kinds[j] != "J":
                return kinds[j]
            j += step
        return None

    keep = []
    for i, (k, t) in enumerate(info):
        if k in ("G",):
            keep.append(t)
        elif k == "P":
            prev_k, next_k = neighbor_kind(i, -1), neighbor_kind(i, 1)
            if prev_k in ("G", "N") and next_k in ("G", "N"):
                keep.append(t)
        elif k == "N":
            p = kinds[i - 1] if i > 0 else None
            n = kinds[i + 1] if i + 1 < len(kinds) else None
            if p == "J" and n == "J":
                continue
            if p == "G" or n == "G" or (p is None and n == "T") or (n is None and p == "T"):
                keep.append(t)
        elif k == "T":
            p = kinds[i - 1] if i > 0 else None
            n = kinds[i + 1] if i + 1 < len(kinds) else None
            if p in ("G", "P", "N") and n in ("G", "P", "N"):
                keep.append(t)
    return " ".join(keep), len(toks), len(toks) - len(keep)


# ============================================================================
# QUESTION EXTRACTION
# ============================================================================
_Q_HEAD = re.compile(r"^\s*(\d{1,3})\s*([.)]|,)\s*(.*)$")
_OPT_MARK = re.compile(r"(?:(?<=\s)|^)\(\s*([A-E])\s*\)")
_HDR_ABCD = re.compile(r"^\s*\(\s*a\s*\)\s*\(\s*b\s*\)\s*\(\s*c\s*\)\s*\(\s*d\s*\)\s*$", re.I)
_MAX_LINES_PER_QUESTION = 40


def _new_q(n):
    return {"n": n, "text": [], "opts": {}, "hdr": False, "nlines": 0, "last": None}


def _add_text(q, text):
    if q["last"] == "_dup":
        return
    if q["last"]:
        q["opts"][q["last"]].append(text)
    else:
        q["text"].append(text)


def parse_questions(lines, lo, hi):
    found, cur = {}, None

    def options_seen(q):
        return len(q["opts"]) >= 2 or "E" in q["opts"]

    for raw in lines:
        line = normalize(raw).strip()
        if not line:
            continue
        m = _Q_HEAD.match(line)
        if m:
            n, sep, rest = int(m.group(1)), m.group(2), m.group(3)
            start = False
            if cur is None:
                start = lo <= n <= hi and sep != ","
            elif cur["n"] < n <= cur["n"] + 5:
                if options_seen(cur):
                    start = True
                elif n == cur["n"] + 1 and sep != "," and cur["nlines"] >= 12:
                    start = True
            if start:
                cur = _new_q(n)
                if lo <= n <= hi and n not in found:
                    found[n] = cur
                line = rest.strip()
                if not line:
                    continue
        if cur is None:
            continue
        cur["nlines"] += 1
        if cur["nlines"] > _MAX_LINES_PER_QUESTION:
            continue
        if _HDR_ABCD.match(line) and not cur["opts"]:
            cur["hdr"] = True
            continue
        marks = list(_OPT_MARK.finditer(line))
        if not marks:
            _add_text(cur, line)
            continue
        pre = line[: marks[0].start()].strip()
        if pre:
            _add_text(cur, pre)
        for k, mk in enumerate(marks):
            end = marks[k + 1].start() if k + 1 < len(marks) else len(line)
            label, body = mk.group(1), line[mk.end(): end].strip()
            if label in cur["opts"]:
                cur["last"] = "_dup"
            else:
                cur["opts"][label] = [body] if body else []
                cur["last"] = label
    return found


def _tidy(text):
    text = re.sub(r"\s+([,.;:?!])", r"\1", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def finalize_question(q):
    removed = total = 0
    qlines = []
    for ln in q["text"]:
        c, a, r = clean_line(ln, False)
        total, removed = total + a, removed + r
        if c:
            qlines.append(c)
    question = _tidy("\n".join(qlines)).replace(" \n", "\n")

    out = ["(a) (b) (c) (d)"] if q["hdr"] else []
    for label in "ABCDE":
        if label not in q["opts"]:
            continue
        body = " ".join(p for p in q["opts"][label] if p)
        c, a, r = clean_line(body, True)
        total, removed = total + a, removed + r
        c = _tidy(c)
        if label == "E":
            c = E_CANON
        out.append(f"({label}){c}")
    return question, "\n".join(out), set(q["opts"]), (removed / total if total else 0.0)


def extract_questions(lines, lo=1, hi=100):
    found = parse_questions(lines, lo, hi)
    rows, review = [], []
    for n in range(lo, hi + 1):
        if n not in found:
            review.append((n, "Question number not found in the PDF text"))
            continue
        question, answers, labels, junk = finalize_question(found[n])
        rows.append({"number": n, "question": question, "answers": answers})
        problems = []
        if len(question) < 8:
            problems.append("question text missing or too short")
        elif junk >= 0.25:
            problems.append("a lot of junk was removed - verify the wording")
        miss = [x for x in "ABCD" if x not in labels]
        if miss:
            problems.append("options missing: " + ", ".join(miss))
        if problems:
            review.append((n, "; ".join(problems)))
    return rows, review


def save_questions_xlsx(rows, review, raw_pages, path, series="", bank=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(["TestSeries", "QuestionBank", "QuestionNo", "Number", "Question", "Anwers", "CorrectAnswer"])
    bank_val = int(bank) if str(bank).strip().isdigit() else (bank or None)
    for row in rows:
        r = ws.max_row + 1
        ws.append([series or None, bank_val, f'=CONCATENATE(A{r},"-",B{r},"-",D{r})',
                   row["number"], row["question"], row["answers"], None])
    for col, width in zip("ABCDEFG", (12, 14, 18, 9, 70, 70, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    for row_cells in ws.iter_rows(min_row=2):
        for c in row_cells:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    rv = wb.create_sheet("Needs review")
    rv.append(["Number", "Problem"])
    for n, p in sorted(review):
        rv.append([n, p])
    rv.column_dimensions["B"].width = 100

    raw = wb.create_sheet("Raw text")
    raw.append(["Page", "Column", "Line", "Text"])
    for page_no, words in raw_pages.items():
        for line_no, col, text, _ in group_lines(words):
            raw.append([page_no, col, line_no, text])
    raw.column_dimensions["D"].width = 100

    font_tamil = Font(name="Noto Sans Tamil", size=10)
    font_tamil_bold = Font(name="Noto Sans Tamil", size=10, bold=True)
    for sheet in (ws, rv, raw):
        for c in sheet[1]:
            c.font = font_tamil_bold
        for row_cells in sheet.iter_rows(min_row=2):
            for c in row_cells:
                c.font = font_tamil
    wb.save(path)


def save_questions_txt(rows, review, raw_pages, path, series="", bank=""):
    with open(path, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(f"{row['number']}. {row['question']}\n{row['answers']}\n\n")


def save_questions_docx(rows, review, raw_pages, path, series="", bank=""):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    try:
        style = doc.styles['Normal']
        style.font.name = 'Noto Sans Tamil'
        style.font.size = Pt(11)
    except Exception:
        pass

    for row in rows:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{row['number']}. {row['question']}")
        r1.font.name = "Noto Sans Tamil"
        r1.bold = True
        for line in row["answers"].split("\n"):
            p2 = doc.add_paragraph()
            r2 = p2.add_run(line)
            r2.font.name = "Noto Sans Tamil"
    doc.save(path)


QUESTION_SAVERS = {"xlsx": save_questions_xlsx, "docx": save_questions_docx, "txt": save_questions_txt}


# ============================================================================
# WINDOW (Tkinter)
# ============================================================================
import tkinter as tk
from tkinter import filedialog, messagebox, ttk


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Tamil PDF Tool - batch OCR + questions 1 to 100")
        self.geometry("820x790")
        self.minsize(700, 640)
        try:
            self.option_add("*Font", ("Noto Sans Tamil", 10))
        except Exception:
            pass

        self.msgs = queue.Queue()
        self.stop_flag = threading.Event()
        self.worker = None
        self.chosen_files = []

        self.folder = tk.StringVar()
        self.fmt_xlsx = tk.BooleanVar(value=True)
        self.fmt_docx = tk.BooleanVar(value=False)
        self.fmt_txt = tk.BooleanVar(value=True)
        self.recursive = tk.BooleanVar(value=False)
        self.overwrite = tk.BooleanVar(value=False)
        self.extract = tk.BooleanVar(value=True)
        self.q_from = tk.IntVar(value=1)
        self.q_to = tk.IntVar(value=100)
        self.series = tk.StringVar(value="TNPSC")
        self.bank = tk.StringVar(value="20242")
        self.batch = tk.IntVar(value=32)
        self.page_batch = tk.IntVar(value=4)
        self.dpi = tk.IntVar(value=150)
        self.min_conf = tk.DoubleVar(value=0.85)
        self.lang = tk.StringVar(value="tamil,english")
        self.columns = tk.BooleanVar(value=True)
        self.use_gpu = tk.BooleanVar(value=True)
        self.straight = tk.BooleanVar(value=True)

        self._build()
        self.after(100, self._drain)

    def _build(self):
        pad = {"padx": 10, "pady": 5}

        top = ttk.LabelFrame(self, text="1. Input: select PDF file(s) or a whole folder")
        top.pack(fill="x", **pad)
        ttk.Entry(top, textvariable=self.folder).pack(side="left", fill="x", expand=True, padx=8, pady=8)
        ttk.Button(top, text="Select PDF files...", command=self.pick_files).pack(side="left", padx=4)
        ttk.Button(top, text="Select folder...", command=self.pick_folder).pack(side="left", padx=(4, 8))

        qf = ttk.LabelFrame(self, text="2. What to extract")
        qf.pack(fill="x", **pad)
        r1 = ttk.Frame(qf)
        r1.pack(fill="x", padx=8, pady=4)
        ttk.Checkbutton(r1, text="Extract questions numbered", variable=self.extract).pack(side="left")
        ttk.Spinbox(r1, from_=1, to=999, textvariable=self.q_from, width=5).pack(side="left", padx=6)
        ttk.Label(r1, text="to").pack(side="left")
        ttk.Spinbox(r1, from_=1, to=999, textvariable=self.q_to, width=5).pack(side="left", padx=6)
        ttk.Label(r1, text="  Test series:").pack(side="left")
        ttk.Entry(r1, textvariable=self.series, width=10).pack(side="left", padx=4)
        ttk.Label(r1, text="Question bank:").pack(side="left")
        ttk.Entry(r1, textvariable=self.bank, width=10).pack(side="left", padx=4)
        ttk.Label(
            qf,
            text="On: Question column = question text only, Anwers column = options (A)-(E) only; other numbers are dropped.\n"
                 "Off: the plain page text is saved instead. The correct-answer column is left blank (a question paper has no key).",
            foreground="#555", justify="left",
        ).pack(anchor="w", padx=8, pady=(0, 6))

        opt = ttk.LabelFrame(self, text="3. Options")
        opt.pack(fill="x", **pad)

        row1 = ttk.Frame(opt)
        row1.pack(fill="x", padx=8, pady=4)
        ttk.Label(row1, text="Save as:").pack(side="left")
        ttk.Checkbutton(row1, text="Excel (.xlsx)", variable=self.fmt_xlsx).pack(side="left", padx=6)
        ttk.Checkbutton(row1, text="Word (.docx)", variable=self.fmt_docx).pack(side="left", padx=6)
        ttk.Checkbutton(row1, text="Text (.txt)", variable=self.fmt_txt).pack(side="left", padx=6)

        row2 = ttk.Frame(opt)
        row2.pack(fill="x", padx=8, pady=4)
        ttk.Label(row2, text="Batch size (words per run):").pack(side="left")
        ttk.Spinbox(row2, from_=1, to=1024, textvariable=self.batch, width=6).pack(side="left", padx=(4, 14))
        ttk.Label(row2, text="Pages per batch:").pack(side="left")
        ttk.Spinbox(row2, from_=1, to=50, textvariable=self.page_batch, width=5).pack(side="left", padx=(4, 14))
        ttk.Label(row2, text="PDF DPI:").pack(side="left")
        ttk.Spinbox(row2, from_=72, to=600, increment=25, textvariable=self.dpi, width=6).pack(side="left", padx=4)

        row3 = ttk.Frame(opt)
        row3.pack(fill="x", padx=8, pady=4)
        ttk.Label(row3, text="Language:").pack(side="left")
        ttk.Combobox(row3, textvariable=self.lang, width=14, state="readonly",
                     values=["tamil,english", "tamil", "english"]).pack(side="left", padx=(4, 14))
        ttk.Label(row3, text="Min word confidence:").pack(side="left")
        ttk.Spinbox(row3, from_=0.10, to=0.99, increment=0.05, format="%.2f",
                    textvariable=self.min_conf, width=6).pack(side="left", padx=4)

        row4 = ttk.Frame(opt)
        row4.pack(fill="x", padx=8, pady=4)
        ttk.Checkbutton(row4, text="Detect columns (reading order)", variable=self.columns).pack(side="left", padx=(0, 14))
        ttk.Checkbutton(row4, text="Straight pages (faster)", variable=self.straight).pack(side="left", padx=(0, 14))
        ttk.Checkbutton(row4, text="Use GPU if available", variable=self.use_gpu).pack(side="left")

        row5 = ttk.Frame(opt)
        row5.pack(fill="x", padx=8, pady=4)
        ttk.Checkbutton(row5, text="Include sub-folders", variable=self.recursive).pack(side="left", padx=(0, 16))
        ttk.Checkbutton(row5, text="Redo files that are already converted", variable=self.overwrite).pack(side="left")

        run = ttk.LabelFrame(self, text="4. Convert")
        run.pack(fill="both", expand=True, **pad)
        btns = ttk.Frame(run)
        btns.pack(fill="x", padx=8, pady=6)
        self.start_btn = ttk.Button(btns, text="Start", command=self.start)
        self.start_btn.pack(side="left")
        self.stop_btn = ttk.Button(btns, text="Stop", command=self.stop, state="disabled")
        self.stop_btn.pack(side="left", padx=8)
        self.status = ttk.Label(btns, text="Ready")
        self.status.pack(side="left", padx=12)

        self.bar = ttk.Progressbar(run, mode="determinate")
        self.bar.pack(fill="x", padx=8)
        logf = ttk.Frame(run)
        logf.pack(fill="both", expand=True, padx=8, pady=8)
        self.log = tk.Text(logf, height=8, wrap="word", font=("Noto Sans Tamil", 10), state="disabled")
        sb = ttk.Scrollbar(logf, command=self.log.yview)
        self.log.configure(yscrollcommand=sb.set)
        self.log.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

    def pick_folder(self):
        d = filedialog.askdirectory(title="Select folder with PDFs / images")
        if d:
            self.chosen_files = []
            self.folder.set(d)

    def pick_files(self):
        fs = filedialog.askopenfilenames(
            title="Select PDF file(s)",
            filetypes=[("PDF files", "*.pdf"), ("PDF and images", "*.pdf *.png *.jpg *.jpeg *.tif *.tiff *.bmp"),
                       ("All files", "*.*")],
        )
        if fs:
            self.chosen_files = [Path(f) for f in fs]
            self.folder.set("; ".join(fs) if len(fs) <= 2 else f"{len(fs)} files selected")

    def write(self, text):
        self.msgs.put(("log", text))

    def start(self):
        folder = self.folder.get().strip()
        if not self.chosen_files and (not folder or not Path(folder).is_dir()):
            messagebox.showwarning("Select input", "Select PDF file(s) or a folder first.")
            return
        formats = [f for f, v in (("xlsx", self.fmt_xlsx), ("docx", self.fmt_docx), ("txt", self.fmt_txt)) if v.get()]
        if not formats:
            messagebox.showwarning("Select format", "Tick at least one output format.")
            return
        try:
            batch, page_batch, dpi = int(self.batch.get()), int(self.page_batch.get()), int(self.dpi.get())
            min_conf = float(self.min_conf.get())
            q_from, q_to = int(self.q_from.get()), int(self.q_to.get())
            if batch < 1 or page_batch < 1 or dpi < 36 or not (0.0 < min_conf < 1.0) or q_from < 1 or q_to < q_from:
                raise ValueError
        except Exception:
            messagebox.showwarning("Check numbers",
                                   "Batch size, pages per batch and DPI must be valid numbers, min confidence must be "
                                   "between 0.10 and 0.99, and the question range must be like 1 to 100.")
            return

        self.stop_flag.clear()
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.bar["value"] = 0
        self.log.config(state="normal")
        self.log.delete("1.0", "end")
        self.log.config(state="disabled")

        job = dict(
            folder=folder, files=list(self.chosen_files), formats=formats, batch=batch, page_batch=page_batch,
            dpi=dpi, min_conf=min_conf, langs=[l.strip() for l in self.lang.get().split(",") if l.strip()],
            recursive=self.recursive.get(), overwrite=self.overwrite.get(), columns=self.columns.get(),
            use_gpu=self.use_gpu.get(), straight=self.straight.get(), extract=self.extract.get(),
            q_from=q_from, q_to=q_to, series=self.series.get().strip(), bank=self.bank.get().strip(),
        )
        self.worker = threading.Thread(target=self.run_job, args=(job,), daemon=True)
        self.worker.start()

    def stop(self):
        self.stop_flag.set()
        self.stop_btn.config(state="disabled")
        self.status.config(text="Stopping after the current batch of pages...")

    def run_job(self, a):
        try:
            self._run_job(a)
        except Exception as e:
            self.write(f"\nStopped because of an error: {e}")
            self.msgs.put(("done", None))

    def _save_outputs(self, a, results, outputs):
        if a["extract"]:
            lines = [t for p in sorted(results) for _, _, t, _ in group_lines(results[p])]
            rows, review = extract_questions(lines, a["q_from"], a["q_to"])
            for fmt_name, out_path in outputs.items():
                QUESTION_SAVERS[fmt_name](rows, review, results, str(out_path), a["series"], a["bank"])
                self.write(f"    saved {out_path}")
            bad = [r for r in review if "not found" not in r[1]]
            missing = [r[0] for r in review if "not found" in r[1]]
            self.write(f"    questions found: {len(rows)} of {a['q_to'] - a['q_from'] + 1}; "
                       f"{len(bad)} need review" + (f"; not found: {_compact(missing)}" if missing else ""))
        else:
            for fmt_name, out_path in outputs.items():
                SAVERS[fmt_name](results, str(out_path))
                self.write(f"    saved {out_path}")

    def _run_job(self, a):
        if a.get("files"):
            files = [Path(f) for f in a["files"]]
        else:
            try:
                files = collect_inputs(a["folder"], a["recursive"])
            except SystemExit as e:
                self.write(str(e))
                self.msgs.put(("done", None))
                return
        if not files:
            self.write("No PDF or image files found in that folder.")
            self.msgs.put(("done", None))
            return

        plan, skipped, failed, done = [], 0, [], 0
        for src in files:
            stem = output_stem(src, files)
            outputs = {f: Path(f"{stem}.{f}") for f in a["formats"]}
            if not a["overwrite"] and all(p.exists() for p in outputs.values()):
                self.write(f"skipped (already converted): {src.name}")
                skipped += 1
                continue
            plan.append((src, outputs, count_pages(src)))
        total_pages = sum(n for _, _, n in plan)
        if not plan:
            self.write(f"\nNothing to do: {skipped} file(s) already converted.")
            self.msgs.put(("done", None))
            return

        self.msgs.put(("max", total_pages))
        self.write(f"{len(plan)} file(s), {total_pages} page(s) to read. Loading OCR model ...")
        self.msgs.put(("status", "Loading model..."))
        try:
            ocr = load_engine(a["langs"], a["batch"], a["min_conf"], a["use_gpu"], a["straight"])
        except Exception as e:
            self.write(f"Could not load the OCR model: {e}")
            self.msgs.put(("done", None))
            return
        device = engine_device(ocr)
        self.write(f"Running on: {device}")

        t_start, counter = time.time(), {"pages": 0}
        for n, (src, outputs, npages) in enumerate(plan, 1):
            if self.stop_flag.is_set():
                break
            self.write(f"[{n}/{len(plan)}] {src}  ({npages} page(s))")
            file_start = counter["pages"]

            def progress(info, n=n):
                counter["pages"] += len(info["pages"])
                got = counter["pages"]
                avg = (time.time() - t_start) / got
                p = info["pages"]
                label = f"{p[0]}" if len(p) == 1 else f"{p[0]}-{p[-1]}"
                self.write(f"    pages {label}: {info['seconds']:.1f}s  "
                           f"(find words {info['detect']:.1f}s, read {info['read']:.1f}s, {info['words']} words)")
                self.msgs.put(("step", got))
                self.msgs.put(("status", f"File {n}/{len(plan)} - page {got}/{total_pages} - "
                                         f"avg {avg:.1f}s/page - about {fmt_time(avg * (total_pages - got))} left"))

            try:
                results, was_stopped = ocr_document(ocr, src, a["dpi"], None, a["page_batch"], a["columns"],
                                                    should_stop=self.stop_flag.is_set, on_progress=progress)
                if was_stopped:
                    self.write("Stopped by user (this file was not saved).")
                    break
                self._save_outputs(a, results, outputs)
                done += 1
            except Exception as e:
                self.write(f"    FAILED: {e}")
                failed.append(src.name)
                counter["pages"] = file_start + npages
                self.msgs.put(("step", counter["pages"]))

        self.write(f"\nFinished in {fmt_time(time.time() - t_start)}: "
                   f"{done} converted, {skipped} skipped, {len(failed)} failed.")
        self.msgs.put(("done", None))

    def _drain(self):
        try:
            while True:
                kind, val = self.msgs.get_nowait()
                if kind == "log":
                    self.log.config(state="normal")
                    self.log.insert("end", val + "\n")
                    self.log.see("end")
                    self.log.config(state="disabled")
                elif kind == "max":
                    self.bar["maximum"] = val
                elif kind == "step":
                    self.bar["value"] = val
                elif kind == "status":
                    self.status.config(text=val)
                elif kind == "done":
                    self.start_btn.config(state="normal")
                    self.stop_btn.config(state="disabled")
                    self.status.config(text="Done")
        except queue.Empty:
            pass
        self.after(100, self._drain)


def _compact(nums, limit=12):
    nums = list(nums)
    s = ", ".join(str(n) for n in nums[:limit])
    return s + (f" ... (+{len(nums) - limit} more)" if len(nums) > limit else "")


if __name__ == "__main__":
    App().mainloop()
